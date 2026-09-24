"""Die mitgelieferten Standard-Vorlagen müssen mit dem Kontext vollständig befüllbar sein."""

from docx import Document

from bewerbungdb import docx_engine, templates
from bewerbungdb.config import Settings
from bewerbungdb.context import build_context

JOB = {"title": "Tester", "company": "ACME GmbH", "contact_last_name": "Muster", "contact_street": "Weg",
       "contact_street_nr": "1", "contact_plz": "70000", "contact_city": "Stadt", "bewerbungstext": "Text."}
PROFIL = {"vorname": "Max", "nachname": "Mustermann", "email": "m@x.de", "telefon": "1", "strasse": "S",
          "hausnummer": "2", "plz": "70000", "ort": "Stadt"}


def test_defaults_are_copied_once(tmp_path):
    s = Settings(home=tmp_path, config={"vorlagen_pfad": str(tmp_path / "V")}, profil=PROFIL)
    templates.ensure_defaults(s)
    assert {p.name for p in (tmp_path / "V").glob("blprnt*")} >= {
        "blprnt_Deckblatt.docx", "blprnt_Anschreiben.docx", "blprnt_Lebenslauf.docx", "blprnt_Email.txt"}
    marker = tmp_path / "V" / "blprnt_Deckblatt.docx"
    marker.write_bytes(b"eigene Vorlage")
    templates.ensure_defaults(s)  # darf vorhandene Vorlagen nicht überschreiben
    assert marker.read_bytes() == b"eigene Vorlage"


def test_default_templates_fill_completely(tmp_path):
    ctx = build_context(JOB, PROFIL)
    for name in ("Deckblatt", "Anschreiben", "Lebenslauf"):
        src = templates.DEFAULTS_DIR / f"blprnt_{name}.docx"
        assert docx_engine.missing_placeholders(src, ctx) == []
        out = docx_engine.fill_docx(src, ctx, tmp_path / f"{name}.docx")
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "{{" not in text and "Mustermann" in text
