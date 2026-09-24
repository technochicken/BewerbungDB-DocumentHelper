"""Erzeugt neutrale Standard-Vorlagen (ohne persönliche Daten) in bewerbungdb/default_templates.

Ausführen: python tools/make_default_templates.py
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / "bewerbungdb" / "default_templates"
BLUE = RGBColor(0x1F, 0x38, 0x64)


def base_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    st.paragraph_format.space_after = Pt(0)
    return doc


def para(doc, text="", size=None, bold=False, color=None, align=None, after=0, before=0):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
    return p


def deckblatt():
    d = base_doc()
    para(d, "{{ name }}", 11, True, BLUE, WD_ALIGN_PARAGRAPH.RIGHT, after=150)
    para(d, "BEWERBUNG", 44, True, BLUE, after=20)
    para(d, "als {{ bereinigter_jobtitel }}", 20, True, after=12)
    para(d, "{{ firmen_floskel }}", 12, False, after=0)
    d.save(OUT / "blprnt_Deckblatt.docx")


def anschreiben():
    d = base_doc()
    para(d, "{{ name }}", 14, True, BLUE, WD_ALIGN_PARAGRAPH.RIGHT)
    para(d, "{{ profil_adresse }}", 9, align=WD_ALIGN_PARAGRAPH.RIGHT, after=36)
    para(d, "{{ firma }}")
    para(d, "{{ ansprechpartner }}")
    para(d, "{{ stelle_adresse }}")
    para(d, "{{ stelle_plz }} {{ stelle_ort }}", after=24)
    para(d, "{{ profil_ort }}, {{ datum }}", align=WD_ALIGN_PARAGRAPH.RIGHT, after=14)
    para(d, "Bewerbung als {{ bereinigter_jobtitel }}", 12, True, after=14)
    para(d, "{{ anrede }}", after=10)
    para(d, "{{ anschreiben }}", after=14)
    para(d, "Mit freundlichen Grüßen", after=24)
    para(d, "{{ name }}")
    d.save(OUT / "blprnt_Anschreiben.docx")


def lebenslauf():
    d = base_doc()
    para(d, "{{ name }}", 22, True, BLUE, after=4)
    para(d, "{{ profil_adresse }}")
    para(d, "{{ email }}  ·  {{ telefon }}", after=18)
    para(d, "Lebenslauf", 14, True, BLUE, after=8)
    para(d, "Trage hier deine Stationen ein (Berufserfahrung, Ausbildung, Kenntnisse). "
            "Diese Vorlage kannst du in Word beliebig anpassen; Platzhalter in "
            "doppelten geschweiften Klammern werden automatisch ersetzt.")
    d.save(OUT / "blprnt_Lebenslauf.docx")


def email():
    (OUT / "blprnt_Email.txt").write_text(
        "Betreff: Bewerbung als {{bereinigter_jobtitel}}\n\n"
        "{{anrede}}\n\n"
        "anbei erhalten Sie meine Bewerbungsunterlagen für die Position als {{bereinigter_jobtitel}}.\n\n"
        "Für Rückfragen stehe ich Ihnen gerne zur Verfügung und freue mich auf Ihre Rückmeldung.\n\n"
        "Mit freundlichen Grüßen\n\n{{name}}\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    deckblatt(); anschreiben(); lebenslauf(); email()
    print("Vorlagen erzeugt in", OUT)
