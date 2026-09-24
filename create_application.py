#!/usr/bin/env python3
"""
Standalone Bewerbungsdokument-Generator
Liest Job-Daten von einer REST-API und befüllt DOCX-Vorlagen.

Aufruf:
    python create_application.py [arbeitsordner]

Wenn kein Ordner angegeben wird, gilt das aktuelle Verzeichnis als Arbeitsordner.
"""

from __future__ import annotations

import html
import json
import re
import shutil
import subprocess
import sys
import tempfile
import datetime
from copy import deepcopy
from pathlib import Path
from typing import Any, Optional

# ─────────────────────────────────────────────────────────────────────────────
# Paths
# ─────────────────────────────────────────────────────────────────────────────

SCRIPT_DIR = Path(__file__).parent.resolve()
CONFIG_FILE = SCRIPT_DIR / "config.json"
PROFIL_FILE = SCRIPT_DIR / "profil.json"

DEFAULT_CONFIG: dict = {
    "api_base_url": "https://dev.mach-tec.de",
    "api_key": "",
    "libreoffice_path": "C:/Program Files/LibreOffice/program/soffice.exe",
    "vorlagen_pfad": "../Vorlagen",
}

DEFAULT_PROFIL: dict = {
    "vorname": "Vorname",
    "nachname": "Nachname",
    "email": "email@example.com",
    "telefon": "",
    "strasse": "Musterstraße",
    "hausnummer": "1",
    "plz": "12345",
    "ort": "Musterstadt",
}


# ─────────────────────────────────────────────────────────────────────────────
# Config & Profile
# ─────────────────────────────────────────────────────────────────────────────

def load_config() -> dict:
    if not CONFIG_FILE.exists():
        CONFIG_FILE.write_text(
            json.dumps(DEFAULT_CONFIG, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return dict(DEFAULT_CONFIG)
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise RuntimeError(f"Fehler in config.json: {e}")


def load_profil() -> dict:
    if not PROFIL_FILE.exists():
        PROFIL_FILE.write_text(
            json.dumps(DEFAULT_PROFIL, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return dict(DEFAULT_PROFIL)
    try:
        return json.loads(PROFIL_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise RuntimeError(f"Fehler in profil.json: {e}")


# ─────────────────────────────────────────────────────────────────────────────
# Job-ID detection
# ─────────────────────────────────────────────────────────────────────────────

def find_job_id_file(work_dir: Path) -> Optional[str]:
    """Return the stem of the first .txt file in work_dir (treated as job ID)."""
    for txt in sorted(work_dir.glob("*.txt")):
        stem = txt.stem.strip()
        if stem:
            return stem
    return None


# ─────────────────────────────────────────────────────────────────────────────
# API
# ─────────────────────────────────────────────────────────────────────────────

def fetch_job(api_base: str, job_id: str, api_key: str = "") -> dict:
    try:
        import httpx
    except ImportError:
        raise ImportError(
            "httpx nicht installiert.\nBitte ausführen: pip install httpx"
        )
    url = f"{api_base.rstrip('/')}/api/v1/jobs/{job_id}"
    headers = {"X-Api-Key": api_key} if api_key else {}
    with httpx.Client(timeout=30, headers=headers) as client:
        resp = client.get(url)
    if resp.status_code == 401:
        raise RuntimeError("API-Key ungültig oder fehlend (HTTP 401).\nBitte api_key in config.json setzen.")
    if resp.status_code == 404:
        raise RuntimeError(f"Stelle #{job_id} nicht gefunden (HTTP 404).")
    resp.raise_for_status()
    return resp.json()


# ─────────────────────────────────────────────────────────────────────────────
# Context building
# ─────────────────────────────────────────────────────────────────────────────

def _format_datum(raw: Optional[str]) -> str:
    if not raw:
        return datetime.date.today().strftime("%d.%m.%Y")
    try:
        d = datetime.date.fromisoformat(str(raw))
        return d.strftime("%d.%m.%Y")
    except (ValueError, TypeError):
        return str(raw)


def build_context(job: dict, profil: dict) -> dict:
    # ── Job / Stelle ──────────────────────────────────────────────────────────
    title      = job.get("title") or ""
    company    = job.get("company") or ""
    salutation = job.get("contact_salutation") or ""
    c_title    = job.get("contact_title") or ""
    first_name = job.get("contact_first_name") or ""
    last_name  = job.get("contact_last_name") or ""
    street     = job.get("contact_street") or ""
    street_nr  = job.get("contact_street_nr") or ""
    plz        = job.get("contact_plz") or ""
    city       = job.get("contact_city") or ""

    stelle_adresse = f"{street} {street_nr}".strip()
    stelle_adresse_komplett = (
        f"{stelle_adresse}\n{plz} {city}".strip()
        if (plz or city) else stelle_adresse
    )

    # Contact person
    ap_parts = [p for p in [salutation, c_title, last_name] if p]
    ansprechpartner_full = " ".join(ap_parts)

    if salutation == "Frau":
        adressierungsfloskel = "Sehr geehrte"
    elif salutation == "Herr":
        adressierungsfloskel = "Sehr geehrter"
    else:
        adressierungsfloskel = "Sehr geehrte"

    if last_name:
        anrede_titel = f"{c_title} {last_name}".strip() if c_title else last_name
        if salutation:
            anrede = f"{adressierungsfloskel} {salutation} {anrede_titel},"
        else:
            anrede = f"{adressierungsfloskel} {anrede_titel},"
    else:
        anrede = "Sehr geehrte Damen und Herren,"

    # ── Profil ────────────────────────────────────────────────────────────────
    vorname  = profil.get("vorname") or ""
    nachname = profil.get("nachname") or ""
    p_titel  = profil.get("titel") or ""
    name = " ".join(p for p in [p_titel, vorname, nachname] if p)

    p_strasse    = profil.get("strasse") or ""
    p_hausnummer = profil.get("hausnummer") or ""
    p_plz        = profil.get("plz") or ""
    p_ort        = profil.get("ort") or ""
    profil_adresse = f"{p_strasse} {p_hausnummer}, {p_plz} {p_ort}".strip(", ")

    skills_raw         = profil.get("skills") or []
    sprachen_raw       = profil.get("sprachen") or []
    ausbildung_raw     = profil.get("ausbildung") or []
    berufserfahrung_raw = profil.get("berufserfahrung") or []
    zertifikate_raw    = profil.get("zertifikate") or []
    hobbys_raw         = profil.get("hobbys") or []

    skills_text = ", ".join(s.get("name", "") for s in skills_raw if s.get("name"))
    sprachen_text = ", ".join(
        f"{s.get('sprache', '')} ({s.get('niveau', '')})"
        for s in sprachen_raw if s.get("sprache")
    )
    hobbys_text = ", ".join(
        h if isinstance(h, str) else h.get("name", "") for h in hobbys_raw
    )

    return {
        # Stelle
        "jobtitel":                title,
        "arbeitgeber":             company,
        "bereinigter_jobtitel":    job.get("job_name_personalized") or title,
        "firma":                   company,
        "firmen_floskel":          job.get("company_floskel") or f"bei {company}",
        # Adresse Stelle
        "stelle_strasse":          street,
        "stelle_hausnummer":       street_nr,
        "stelle_plz":              plz,
        "stelle_ort":              city,
        "stelle_adresse":          stelle_adresse,
        "stelle_adresse_komplett": stelle_adresse_komplett,
        # Ansprechpartner
        "ansprechpartner":             ansprechpartner_full,
        "ansprechpartner_formatiert":  last_name,
        "ansprechpartner_vorname":     first_name,
        "ansprechpartner_nachname":    last_name,
        "ansprechpartner_anrede":      salutation,
        "ansprechpartner_email":       job.get("contact_email") or "",
        "ap_anrede":                   salutation,
        "adressierungsfloskel":        adressierungsfloskel,
        "begruessung_floskel":         adressierungsfloskel,
        "anrede":                      anrede,
        # Bewerbung
        "anschreiben": job.get("bewerbungstext") or "",
        "datum":       _format_datum(job.get("application_date")),
        # Profil – Grunddaten
        "name":                  name,
        "vorname":               vorname,
        "nachname":              nachname,
        "titel":                 p_titel,
        "geburtsdatum":          profil.get("geburtsdatum") or "",
        "nationalitaet":         profil.get("nationalitaet") or "",
        "berufsbezeichnung":     profil.get("berufsbezeichnung") or "",
        "berufserfahrung_jahre": str(profil.get("berufserfahrung_jahre") or ""),
        "zusammenfassung":       profil.get("zusammenfassung") or "",
        # Profil – Kontakt
        "email":    profil.get("email") or "",
        "telefon":  profil.get("telefon") or "",
        "mobil":    profil.get("mobil") or "",
        "linkedin": profil.get("linkedin") or "",
        "xing":     profil.get("xing") or "",
        "website":  profil.get("website") or "",
        "github":   profil.get("github") or "",
        # Profil – Adresse
        "profil_strasse":    p_strasse,
        "profil_hausnummer": p_hausnummer,
        "profil_plz":        p_plz,
        "profil_ort":        p_ort,
        "profil_adresse":    profil_adresse,
        # Profil – Listen (für Jinja2-Schleifen in Vorlagen)
        "skills":           skills_raw,
        "skills_text":      skills_text,
        "sprachen":         sprachen_raw,
        "sprachen_text":    sprachen_text,
        "ausbildung":       ausbildung_raw,
        "berufserfahrung":  berufserfahrung_raw,
        "zertifikate":      zertifikate_raw,
        "hobbys":           hobbys_raw,
        "hobbys_text":      hobbys_text,
    }


# ─────────────────────────────────────────────────────────────────────────────
# DOCX filling
# ─────────────────────────────────────────────────────────────────────────────

def fill_docx(template_path: Path, context: dict, output_path: Path) -> Path:
    try:
        from docxtpl import DocxTemplate
    except ImportError:
        raise ImportError(
            "docxtpl nicht installiert.\nBitte ausführen: pip install docxtpl"
        )

    ctx = dict(context)
    # Join word-wrap newlines within logical paragraphs so only real paragraph
    # breaks (\n\n) produce separate DOCX paragraphs.
    anschreiben = ctx.get("anschreiben", "")
    if isinstance(anschreiben, str) and "\n\n" in anschreiben:
        chunks = re.split(r"\n{2,}", anschreiben)
        ctx["anschreiben"] = "\n\n".join(
            " ".join(ln.strip() for ln in chunk.split("\n") if ln.strip())
            for chunk in chunks
        )

    # docxtpl does not auto-escape, so bare & < > in values produce invalid XML.
    render_ctx = {
        k: html.escape(v, quote=False) if isinstance(v, str) else v
        for k, v in ctx.items()
    }

    tpl = DocxTemplate(str(template_path))
    tpl.render(render_ctx)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    tpl.save(str(output_path))

    try:
        _fix_line_breaks(output_path)
    except Exception:
        pass  # Non-fatal: file is still usable without the post-processing

    return output_path


def _fix_line_breaks(docx_path: Path) -> None:
    """Convert <w:br/> soft breaks inserted by docxtpl into proper paragraph breaks."""
    import lxml.etree as etree
    from docx import Document
    from docx.oxml.ns import qn

    W_P    = qn("w:p")
    W_R    = qn("w:r")
    W_BR   = qn("w:br")
    W_PPR  = qn("w:pPr")
    W_RPR  = qn("w:rPr")
    W_TYPE = qn("w:type")

    def is_soft_br(elem: Any) -> bool:
        return elem.tag == W_BR and elem.get(W_TYPE, "") not in ("page", "column")

    doc = Document(str(docx_path))
    body = doc.element.body
    modified = False

    for p_elem in list(body.iter(W_P)):
        pPr = p_elem.find(W_PPR)
        if not any(is_soft_br(c) for r in p_elem.iter(W_R) for c in r):
            continue

        segments: list[list] = [[]]
        for child in list(p_elem):
            if child.tag == W_PPR:
                continue
            if child.tag != W_R:
                segments[-1].append(deepcopy(child))
                continue
            rPr = child.find(W_RPR)
            pending: list = []
            for rc in child:
                if rc.tag == W_RPR:
                    continue
                if is_soft_br(rc):
                    if pending:
                        new_r = etree.Element(W_R)
                        if rPr is not None:
                            new_r.append(deepcopy(rPr))
                        for elem in pending:
                            new_r.append(deepcopy(elem))
                        segments[-1].append(new_r)
                        pending = []
                    segments.append([])
                else:
                    pending.append(deepcopy(rc))
            if pending:
                new_r = etree.Element(W_R)
                if rPr is not None:
                    new_r.append(deepcopy(rPr))
                for elem in pending:
                    new_r.append(deepcopy(elem))
                segments[-1].append(new_r)

        if len(segments) <= 1:
            continue

        for child in list(p_elem):
            if child.tag != W_PPR:
                p_elem.remove(child)
        for elem in segments[0]:
            p_elem.append(elem)

        parent = p_elem.getparent()
        idx = list(parent).index(p_elem)
        for i, seg in enumerate(segments[1:], 1):
            new_p = etree.Element(W_P)
            if pPr is not None:
                new_p.append(deepcopy(pPr))
            for elem in seg:
                new_p.append(elem)
            parent.insert(idx + i, new_p)

        modified = True

    if modified:
        doc.save(str(docx_path))


# ─────────────────────────────────────────────────────────────────────────────
# PDF
# ─────────────────────────────────────────────────────────────────────────────

def docx_to_pdf(docx_path: Path, output_dir: Path, lo_path: str) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    lo = Path(lo_path)

    if not lo.exists():
        lo_cmd = shutil.which("libreoffice") or shutil.which("soffice")
        if not lo_cmd:
            raise FileNotFoundError(
                f"LibreOffice nicht gefunden: {lo_path}\n"
                "Bitte LibreOffice installieren oder den Pfad in config.json korrigieren."
            )
        lo = Path(lo_cmd)

    cmd = [
        str(lo), "--headless", "--convert-to", "pdf",
        "--outdir", str(output_dir), str(docx_path),
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
    if result.returncode != 0:
        raise RuntimeError(f"PDF-Konvertierung fehlgeschlagen:\n{result.stderr}")

    pdf = output_dir / (docx_path.stem + ".pdf")
    if not pdf.exists():
        raise FileNotFoundError(f"Konvertierte PDF nicht gefunden: {pdf}")
    return pdf


def merge_pdfs(pdf_paths: list[Path], output_path: Path) -> Path:
    existing = [p for p in pdf_paths if p.exists()]
    if not existing:
        raise FileNotFoundError("Keine PDF-Dateien zum Zusammenführen gefunden.")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    try:
        import pikepdf
        with pikepdf.Pdf.new() as out:
            for p in existing:
                with pikepdf.Pdf.open(p) as src:
                    out.pages.extend(src.pages)
            out.save(str(output_path))
        return output_path
    except ImportError:
        pass

    try:
        from PyPDF2 import PdfMerger
        merger = PdfMerger()
        for p in existing:
            merger.append(str(p))
        with open(output_path, "wb") as f:
            merger.write(f)
        merger.close()
        return output_path
    except ImportError:
        raise ImportError(
            "Weder pikepdf noch PyPDF2 installiert.\n"
            "Bitte ausführen: pip install pikepdf"
        )


def safe_name(s: str) -> str:
    for src, dst in [("ä", "ae"), ("ö", "oe"), ("ü", "ue"),
                     ("Ä", "Ae"), ("Ö", "Oe"), ("Ü", "Ue"), ("ß", "ss")]:
        s = s.replace(src, dst)
    s = re.sub(r'[\\/:*?"<>|&%$§#*~+\-–]', "", s)
    s = s.replace(" ", "_")
    s = re.sub(r"_+", "_", s)
    return s.strip("_")[:80]


def _fill_text_placeholders(template: str, context: dict) -> str:
    """Replace {{key}} tokens in a plain-text template with context values.

    Unknown or empty keys are left as-is so the user can spot gaps.
    """
    def _repl(m: re.Match) -> str:
        val = context.get(m.group(1).strip())
        return str(val) if val is not None and val != "" else m.group(0)
    return re.sub(r"\{\{(\w+)\}\}", _repl, template)


def _parse_email_template(template: str) -> tuple[Optional[str], str]:
    """Split a blueprint into an optional subject line and body.

    If the first non-empty line starts with 'Betreff:' (case-insensitive),
    that line's value becomes the subject template and the remaining text is
    used as the body.  Otherwise returns (None, template) unchanged.
    """
    lines = template.splitlines()
    for i, line in enumerate(lines):
        if not line.strip():
            continue
        if line.strip().lower().startswith("betreff:"):
            subject = line.strip()[len("betreff:"):].strip()
            body = "\n".join(lines[i + 1:]).lstrip("\n")
            return subject, body
        break
    return None, template


def find_email_template(work_dir: Path, vorlagen_dir: Optional[Path]) -> Optional[str]:
    """Return the content of the first blprnt*Email*.txt found, or None."""
    for d in filter(None, [work_dir, vorlagen_dir, SCRIPT_DIR]):
        for p in sorted(d.glob("blprnt*Email*.txt")):
            try:
                return p.read_text(encoding="utf-8")
            except Exception:
                pass
    return None


def write_eml(
    output_path: Path,
    context: dict,
    absender_email: str,
    pdf_path: Optional[Path] = None,
    body_template: Optional[str] = None,
    subject_template: Optional[str] = None,
) -> Path:
    """Write a ready-to-send .eml draft.

    If *body_template* is provided it is used as the email body after
    {{placeholder}} substitution; otherwise the anschreiben text from
    the context is used verbatim.  If *subject_template* is provided it
    is used as the subject after placeholder substitution; otherwise a
    default subject is derived from jobtitel and firma.
    """
    from email import encoders
    from email.mime.base import MIMEBase
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText

    jobtitel = context.get("bereinigter_jobtitel") or context.get("jobtitel") or "Stelle"
    firma    = context.get("firma") or ""

    if subject_template is not None:
        betreff = _fill_text_placeholders(subject_template, context)
    else:
        betreff = f"Bewerbung als {jobtitel}" + (f" bei {firma}" if firma else "")

    body = (
        _fill_text_placeholders(body_template, context)
        if body_template is not None
        else (context.get("anschreiben") or "")
    )

    msg = MIMEMultipart("mixed")
    msg["Subject"]      = betreff
    msg["MIME-Version"] = "1.0"
    msg["X-Unsent"]     = "1"   # tells Outlook / Thunderbird to open as a compose window
    if absender_email:
        msg["From"] = absender_email
    empfaenger = context.get("ansprechpartner_email") or ""
    if empfaenger:
        msg["To"] = empfaenger

    msg.attach(MIMEText(body, "plain", "utf-8"))

    if pdf_path and pdf_path.exists():
        with open(pdf_path, "rb") as fh:
            att = MIMEBase("application", "pdf")
            att.set_payload(fh.read())
        encoders.encode_base64(att)
        att.add_header("Content-Disposition", "attachment", filename=pdf_path.name)
        msg.attach(att)

    output_path.write_bytes(msg.as_bytes())
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# Blueprint lookup  (work_dir override → vorlagen_dir master → script_dir last)
# ─────────────────────────────────────────────────────────────────────────────

def find_blueprint(typ: str, work_dir: Path, vorlagen_dir: Optional[Path]) -> Optional[Path]:
    pattern = f"blprnt*{typ}*.docx"
    for candidate_dir in filter(None, [work_dir, vorlagen_dir, SCRIPT_DIR]):
        for match in sorted(candidate_dir.glob(pattern)):
            try:
                if get_template_variables(match):  # skip already-rendered files
                    return match
            except Exception:
                pass
    return None


def find_existing_docx(typ: str, work_dir: Path, vorlagen_dir: Optional[Path]) -> Optional[Path]:
    """Find the most recent generated DOCX (Deckblatt_*.docx etc.) in work_dir."""
    matches = sorted(work_dir.glob(f"{typ}_*.docx"), key=lambda p: p.stat().st_mtime)
    if matches:
        return matches[-1]
    return find_blueprint(typ, work_dir, vorlagen_dir)


# ─────────────────────────────────────────────────────────────────────────────
# Placeholder validation
# ─────────────────────────────────────────────────────────────────────────────

# docxtpl injects these names at render time — Jinja2 meta sees them as
# undeclared, but they are always available and should not trigger warnings.
_DOCXTPL_BUILTINS = {
    "RichText", "R", "InlineImage", "Listing", "Subdoc",
    "BlockProtected", "MSO", "hyperlink",
}


def get_template_variables(template_path: Path) -> set[str]:
    """Return all {{ variable }} names found in a DOCX template.

    Strips XML tags before scanning so placeholders split across <w:r> runs
    are rejoined into a single token before the regex runs.
    """
    import zipfile

    variables: set[str] = set()
    with zipfile.ZipFile(str(template_path), "r") as z:
        for entry in z.namelist():
            if not entry.endswith(".xml"):
                continue
            try:
                raw = z.read(entry).decode("utf-8", errors="ignore")
                text = re.sub(r"<[^>]*>", "", raw)
                variables.update(re.findall(r"\{\{-?\s*(\w+)\s*-?\}\}", text))
            except Exception:
                pass
    return variables - _DOCXTPL_BUILTINS


def check_missing_placeholders(
    blueprints: dict[str, Path], context: dict
) -> dict[str, list[str]]:
    """Return {typ: [var, ...]} for {{ var }} placeholders that would render empty.

    A placeholder is flagged when its context value is missing (None) or an
    empty string.  Non-empty lists (loop variables) and zero-values are not
    flagged because they are intentional.
    """
    missing: dict[str, list[str]] = {}
    for typ, path in blueprints.items():
        try:
            tpl_vars = get_template_variables(path)
        except Exception:
            continue
        flagged = []
        for var in sorted(tpl_vars):
            val = context.get(var)
            if val is None or val == "":
                flagged.append(var)
        if flagged:
            missing[typ] = flagged
    return missing


# ─────────────────────────────────────────────────────────────────────────────
# UI (tkinter)
# ─────────────────────────────────────────────────────────────────────────────

def _center(win: Any) -> None:
    win.update_idletasks()
    w, h = win.winfo_width(), win.winfo_height()
    sw, sh = win.winfo_screenwidth(), win.winfo_screenheight()
    win.geometry(f"+{(sw - w) // 2}+{(sh - h) // 2}")


def ask_mode_dialog(root: Any, job: dict) -> Optional[str]:
    import tkinter as tk
    from tkinter import ttk

    result: dict[str, Optional[str]] = {"choice": None}

    dlg = tk.Toplevel(root)
    dlg.title("Bewerbungsdokumente erstellen")
    dlg.resizable(False, False)
    dlg.grab_set()

    pad = {"padx": 20, "pady": 6}

    tk.Label(
        dlg, text="Bewerbungsdokumente erstellen",
        font=("Segoe UI", 11, "bold"),
    ).pack(anchor="w", **pad)

    info_frame = tk.Frame(dlg, bg="#f0f0f0", relief="sunken", bd=1)
    info_frame.pack(fill="x", padx=20, pady=(0, 6))
    tk.Label(
        info_frame,
        text=f"  Stelle:  {job.get('title') or '–'}\n"
             f"  Firma:   {job.get('company') or '–'}\n"
             f"  Ort:     {job.get('location') or '–'}",
        justify="left", bg="#f0f0f0", font=("Segoe UI", 9),
        pady=8, padx=4,
    ).pack(anchor="w")

    tk.Label(dlg, text="Aktion wählen:", font=("Segoe UI", 9, "bold")).pack(
        anchor="w", padx=20, pady=(8, 2)
    )

    mode_var = tk.StringVar(value="new")
    tk.Radiobutton(
        dlg,
        text="Dokumente neu aus Vorlagen erstellen und als PDF exportieren",
        variable=mode_var, value="new",
        font=("Segoe UI", 9),
    ).pack(anchor="w", padx=28)
    tk.Radiobutton(
        dlg,
        text="Nur bestehende DOCX-Dateien als PDF exportieren (kein Überschreiben)",
        variable=mode_var, value="pdf_only",
        font=("Segoe UI", 9),
    ).pack(anchor="w", padx=28, pady=(4, 12))

    btn_frame = tk.Frame(dlg)
    btn_frame.pack(fill="x", padx=20, pady=(0, 16))
    tk.Button(
        btn_frame, text="Abbrechen", width=12,
        command=dlg.destroy,
    ).pack(side="right", padx=(6, 0))
    tk.Button(
        btn_frame, text="Weiter →", width=12,
        command=lambda: [result.__setitem__("choice", mode_var.get()), dlg.destroy()],
        default="active",
    ).pack(side="right")

    dlg.protocol("WM_DELETE_WINDOW", dlg.destroy)
    dlg.update_idletasks()
    dlg.geometry(f"520x{dlg.winfo_reqheight()}")
    _center(dlg)
    root.wait_window(dlg)
    return result["choice"]


def ask_overwrite_dialog(
    root: Any, existing: list[tuple[str, Path]]
) -> Optional[list[str]]:
    """Show checkboxes for existing DOCX files. Returns selected type names, or None if cancelled."""
    import tkinter as tk

    result: dict[str, Optional[list[str]]] = {"choice": None}

    dlg = tk.Toplevel(root)
    dlg.title("Dateien überschreiben?")
    dlg.resizable(False, False)
    dlg.grab_set()

    tk.Label(
        dlg,
        text="Folgende Dateien existieren bereits.\nWelche sollen überschrieben werden?",
        font=("Segoe UI", 9),
        justify="left",
    ).pack(anchor="w", padx=20, pady=(14, 6))

    default_overwrite = {"Deckblatt", "Anschreiben"}
    chk_vars: dict[str, tk.BooleanVar] = {}
    for typ, path in existing:
        var = tk.BooleanVar(value=(typ in default_overwrite))
        chk_vars[typ] = var
        fname = path.name if len(path.name) <= 50 else path.name[:47] + "…"
        tk.Checkbutton(
            dlg,
            text=f"{typ}  ({fname})",
            variable=var,
            font=("Segoe UI", 9),
        ).pack(anchor="w", padx=28, pady=2)

    btn_frame = tk.Frame(dlg)
    btn_frame.pack(fill="x", padx=20, pady=(10, 16))

    def on_ok() -> None:
        result["choice"] = [typ for typ, var in chk_vars.items() if var.get()]
        dlg.destroy()

    tk.Button(btn_frame, text="Abbrechen", width=12, command=dlg.destroy).pack(
        side="right", padx=(6, 0)
    )
    tk.Button(btn_frame, text="Weiter →", width=12, command=on_ok, default="active").pack(
        side="right"
    )

    dlg.protocol("WM_DELETE_WINDOW", dlg.destroy)
    dlg.update_idletasks()
    max_w = dlg.winfo_screenwidth() - 80
    dlg.geometry(f"{min(dlg.winfo_reqwidth(), max_w)}x{dlg.winfo_reqheight()}")
    _center(dlg)
    root.wait_window(dlg)
    return result["choice"]


def show_missing_placeholder_dialog(
    root: Any, missing_by_typ: dict[str, list[str]]
) -> bool:
    """Warn about blueprint placeholders with no matching context key.
    Returns True if user wants to continue anyway, False to abort."""
    import tkinter as tk

    outcome = {"go": False}

    dlg = tk.Toplevel(root)
    dlg.title("Fehlende Angaben")
    dlg.resizable(False, False)
    dlg.grab_set()

    tk.Label(
        dlg,
        text="Folgende Platzhalter sind in den Vorlagen definiert,\n"
             "werden aber von der Datenbank nicht befüllt:",
        font=("Segoe UI", 9),
        fg="#cc6600",
        justify="left",
    ).pack(anchor="w", padx=20, pady=(14, 8))

    for typ, var_list in missing_by_typ.items():
        tk.Label(dlg, text=f"{typ}:", font=("Segoe UI", 9, "bold")).pack(
            anchor="w", padx=20, pady=(4, 0)
        )
        for var in var_list:
            tk.Label(
                dlg,
                text=f"    •  {{{{{var}}}}}",
                font=("Courier New", 9),
                fg="#cc0000",
                justify="left",
            ).pack(anchor="w", padx=28)

    tk.Label(
        dlg,
        text="\nDiese Felder bleiben im Dokument leer.\nTrotzdem fortfahren?",
        font=("Segoe UI", 9),
        justify="left",
    ).pack(anchor="w", padx=20, pady=(6, 4))

    btn_frame = tk.Frame(dlg)
    btn_frame.pack(fill="x", padx=20, pady=(4, 16))

    def on_ok() -> None:
        outcome["go"] = True
        dlg.destroy()

    tk.Button(btn_frame, text="Abbrechen", width=12, command=dlg.destroy).pack(
        side="right", padx=(6, 0)
    )
    tk.Button(
        btn_frame, text="Trotzdem erstellen", width=18, command=on_ok, default="active"
    ).pack(side="right")

    dlg.protocol("WM_DELETE_WINDOW", dlg.destroy)
    dlg.update_idletasks()
    max_w = dlg.winfo_screenwidth() - 80
    dlg.geometry(f"{min(dlg.winfo_reqwidth(), max_w)}x{dlg.winfo_reqheight()}")
    _center(dlg)
    root.wait_window(dlg)
    return outcome["go"]


def show_results_dialog(root: Any, work_dir: Path, results: dict, errors: list[str]) -> None:
    import tkinter as tk

    dlg = tk.Toplevel(root)
    dlg.title("Fertig")
    dlg.resizable(False, False)
    dlg.grab_set()

    def open_folder():
        subprocess.Popen(["explorer", str(work_dir)])

    if errors:
        tk.Label(
            dlg, text=f"{'Fehler' if not results else 'Mit Fehlern abgeschlossen'}:",
            font=("Segoe UI", 10, "bold"),
            fg="#cc0000" if not results else "#cc6600",
        ).pack(anchor="w", padx=20, pady=(14, 0))
        for e in errors:
            tk.Label(
                dlg, text=f"  • {e}", fg="#cc0000",
                wraplength=500, justify="left", font=("Segoe UI", 9),
            ).pack(anchor="w", padx=20)

    if results:
        tk.Label(
            dlg, text="Erstellte Dateien:", font=("Segoe UI", 10, "bold"),
        ).pack(anchor="w", padx=20, pady=(12, 0))
        for label, path in results.items():
            tk.Label(
                dlg,
                text=f"  ✓  {label}:  {path.name}",
                wraplength=500, justify="left", font=("Segoe UI", 9),
                fg="#006600",
            ).pack(anchor="w", padx=20)

    btn_frame = tk.Frame(dlg)
    btn_frame.pack(fill="x", padx=20, pady=16)
    tk.Button(btn_frame, text="Ordner öffnen", width=14, command=open_folder).pack(
        side="left"
    )
    tk.Button(btn_frame, text="Schließen", width=12, command=dlg.destroy).pack(
        side="right"
    )

    dlg.update_idletasks()
    dlg.geometry(f"540x{dlg.winfo_reqheight()}")
    _center(dlg)
    root.wait_window(dlg)


# ─────────────────────────────────────────────────────────────────────────────
# Main
# ─────────────────────────────────────────────────────────────────────────────

TYPES = ["Deckblatt", "Anschreiben", "Lebenslauf"]


def main() -> None:
    import tkinter as tk
    from tkinter import messagebox, simpledialog

    # Determine work directory
    if len(sys.argv) > 1:
        # Strip any stray trailing quote that MSVCRT inserts when CMD passes
        # "%~dp0" with a trailing backslash (the \" is parsed as a literal ").
        arg = sys.argv[1].rstrip('"').rstrip("\\/")
        work_dir = Path(arg).resolve() if arg else Path.cwd().resolve()
    else:
        work_dir = Path.cwd().resolve()

    root = tk.Tk()
    root.withdraw()

    def fatal(msg: str) -> None:
        messagebox.showerror("Fehler", msg, parent=root)
        root.destroy()
        sys.exit(1)

    # Load config + profile
    try:
        config = load_config()
        profil = load_profil()
    except RuntimeError as e:
        fatal(str(e))
        return

    api_url = config.get("api_base_url") or DEFAULT_CONFIG["api_base_url"]
    api_key = config.get("api_key") or ""
    lo_path = config.get("libreoffice_path") or DEFAULT_CONFIG["libreoffice_path"]

    vorlagen_raw = config.get("vorlagen_pfad") or DEFAULT_CONFIG["vorlagen_pfad"]
    vorlagen_dir = (SCRIPT_DIR / vorlagen_raw).resolve()
    if not vorlagen_dir.exists():
        vorlagen_dir = None  # will fall through to work_dir / script_dir

    # Find job ID
    job_id = find_job_id_file(work_dir)
    if not job_id:
        job_id = simpledialog.askstring(
            "Job-ID eingeben",
            f"Keine *.txt-Datei in\n{work_dir}\ngefunden.\n\nBitte Job-ID eingeben:",
            parent=root,
        )
    if not job_id or not job_id.strip():
        fatal("Keine Job-ID angegeben. Abbruch.")
        return
    job_id = job_id.strip()

    # Fetch job data
    try:
        job = fetch_job(api_url, job_id, api_key)
    except Exception as e:
        fatal(f"API-Fehler beim Abrufen von Stelle #{job_id}:\n\n{e}")
        return

    # Mode selection dialog
    mode = ask_mode_dialog(root, job)
    if not mode:
        root.destroy()
        sys.exit(0)

    # Build context
    context = build_context(job, profil)

    # Build output file suffix
    suffix_parts = [
        safe_name(job.get("job_name_personalized") or job.get("title") or "Bewerbung"),
        safe_name(profil.get("vorname") or ""),
        safe_name(profil.get("nachname") or ""),
    ]
    suffix = "_".join(p for p in suffix_parts if p)

    errors: list[str] = []
    result_docx: dict[str, Path] = {}
    results: dict[str, Path] = {}

    if mode == "new":
        # ── 1. Ask which existing files to overwrite ───────────────────────────
        existing_files = [
            (typ, work_dir / f"{typ}_{suffix}.docx")
            for typ in TYPES
            if (work_dir / f"{typ}_{suffix}.docx").exists()
        ]

        skip_types: set[str] = set()
        if existing_files:
            overwrite_types = ask_overwrite_dialog(root, existing_files)
            if overwrite_types is None:
                root.destroy()
                sys.exit(0)
            overwrite_set = set(overwrite_types)
            for typ, path in existing_files:
                if typ not in overwrite_set:
                    result_docx[typ] = path
                    skip_types.add(typ)

        # ── 2. Resolve blueprints for types being (re)generated ───────────────
        blueprints: dict[str, Path] = {}
        for typ in TYPES:
            if typ in skip_types:
                continue
            bp = find_blueprint(typ, work_dir, vorlagen_dir)
            if bp is not None:
                blueprints[typ] = bp

        # ── 3. Warn about placeholders the DB cannot fill ─────────────────────
        missing = check_missing_placeholders(blueprints, context)
        if missing:
            if not show_missing_placeholder_dialog(root, missing):
                root.destroy()
                sys.exit(0)

        # ── 4. Generate documents ──────────────────────────────────────────────
        for typ in TYPES:
            if typ in skip_types:
                continue
            blueprint = blueprints.get(typ)
            if blueprint is None:
                errors.append(f"Vorlage nicht gefunden: blprnt*{typ}*.docx")
                continue
            out_docx = work_dir / f"{typ}_{suffix}.docx"
            try:
                fill_docx(blueprint, context, out_docx)
                result_docx[typ] = out_docx
                results[f"{typ} (DOCX)"] = out_docx
            except Exception as e:
                errors.append(f"{typ}: DOCX-Fehler: {e}")
    else:
        for typ in TYPES:
            docx = find_existing_docx(typ, work_dir, vorlagen_dir)
            if docx:
                result_docx[typ] = docx

    # PDF conversion + merging
    if result_docx:
        temp_dir = Path(tempfile.mkdtemp(prefix="bewerbung_"))
        try:
            pdf_by_type: dict[str, Path] = {}
            for typ, docx_path in result_docx.items():
                try:
                    pdf = docx_to_pdf(docx_path, temp_dir, lo_path)
                    out_pdf = work_dir / f"{typ}_{suffix}.pdf"
                    shutil.copy2(pdf, out_pdf)
                    pdf_by_type[typ] = out_pdf
                    results[f"{typ} (PDF)"] = out_pdf
                except Exception as e:
                    errors.append(f"{typ} PDF: {e}")

            # Deckblatt + Anschreiben combined
            if "Deckblatt" in pdf_by_type and "Anschreiben" in pdf_by_type:
                kombi = work_dir / f"Deckblatt_Anschreiben_{suffix}.pdf"
                try:
                    merge_pdfs([pdf_by_type["Deckblatt"], pdf_by_type["Anschreiben"]], kombi)
                    results["Deckblatt+Anschreiben"] = kombi
                except Exception as e:
                    errors.append(f"Zusammenführen Deckblatt+Anschreiben: {e}")

            # Complete application PDF
            alle = [pdf_by_type[t] for t in TYPES if t in pdf_by_type]
            anhang = work_dir / "Anhang.pdf"
            if not anhang.exists() and vorlagen_dir:
                anhang = vorlagen_dir / "Anhang.pdf"
            if anhang.exists():
                alle.append(anhang)
            if alle:
                vollstaendig = work_dir / f"Bewerbung_{suffix}.pdf"
                try:
                    merge_pdfs(alle, vollstaendig)
                    results["Vollständige Bewerbung"] = vollstaendig
                except Exception as e:
                    errors.append(f"Zusammenführen Bewerbung: {e}")
        finally:
            shutil.rmtree(temp_dir, ignore_errors=True)

    # ── E-Mail Entwurf ────────────────────────────────────────────────────────
    if context.get("anschreiben"):
        eml_path   = work_dir / f"Email_{suffix}.eml"
        attach_pdf = (
            results.get("Vollständige Bewerbung")
            or results.get("Deckblatt+Anschreiben")
        )
        email_tpl  = find_email_template(work_dir, vorlagen_dir)
        email_subject_tpl: Optional[str] = None
        if email_tpl is not None:
            email_subject_tpl, email_tpl = _parse_email_template(email_tpl)
        try:
            write_eml(
                eml_path,
                context,
                profil.get("email") or "",
                attach_pdf,
                email_tpl,
                email_subject_tpl,
            )
            results["E-Mail Entwurf"] = eml_path
        except Exception as e:
            errors.append(f"E-Mail Entwurf: {e}")

    show_results_dialog(root, work_dir, results, errors)
    root.destroy()


if __name__ == "__main__":
    main()
